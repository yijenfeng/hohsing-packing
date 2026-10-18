import openpyxl, sys, re, json
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
def iso(s, t=True):
    m = re.match(r'(\d+)/(\d+)/(\d+)\(.\)\s*(\d+:\d+)?', str(s))
    d = f'{int(m[1]):04d}-{int(m[2]):02d}-{int(m[3]):02d}'
    return d + ('T' + (m[4].zfill(5)) if t and m[4] else '')
hol = []
for r in wb['參數'].iter_rows(values_only=True):
    if re.match(r'\d+/\d+/\d+\(', str(r[0])): hol.append(iso(r[0], False))
orders = []
for r in wb['總覽'].iter_rows(min_row=10, max_row=15, values_only=True):
    orders.append(dict(id=r[1], cust=r[0], ship=iso(r[2], False), due=iso(r[3], False), fg1e=r[4]))
items = []
for r in wb['品項完成日'].iter_rows(min_row=4, values_only=True):
    if r[0] in ('OHM', 'YAZAWA'):
        items.append(dict(order=r[1], item=r[4], loc=r[5], qty=r[6]))
tasks = []
for n in ['芳排程', '智排程', '幫手1排程', '幫手2排程', '幫手3排程']:
    w = n[:-2]
    for r in wb[n].iter_rows(min_row=3, values_only=True):
        if not isinstance(r[0], int): continue
        tasks.append(dict(id=f'{w}-{r[0]}', w=w, kind='lathe' if r[1] == '車台' else 'pack', order=r[3],
                          item=('FG1E' if r[1] == '車台' else str(r[5]).split('#')[0]), seg=str(r[5]), how=r[6],
                          qty=r[7], box=r[8] or 0, pd=r[9], s=iso(r[10]), e=iso(r[11])))
helpers = {}
for r in wb['總覽'].iter_rows(min_row=32, max_row=34, values_only=True):
    a, b = r[1].split('–'); helpers[r[0]] = [iso(a, False), iso(b, False)]
plan = dict(made='2026-10-16', base='截至 10/15 收工；04-6825、04-9317 已包完，04-9326 已包 8,000',
            holidays=hol, orders=orders, items=items, tasks=tasks, helpers=helpers,
            rates={'吊入': 1300, '袋入': 3400, '卡入': 5100, 'box': 17, 'lathe': 14000, 'chg': 15})
json.dump(plan, open(sys.argv[2], 'w'), ensure_ascii=False, separators=(',', ':'))
print(len(tasks), len(items), helpers, hol[:3])
# sanity: pack qty per item vs items
from collections import Counter
c = Counter()
for t in tasks:
    if t['kind'] == 'pack': c[(t['order'], t['item'])] += t['qty']
bad = [(i['order'], i['item'], i['qty'], c.get((i['order'], i['item']))) for i in items if i['loc'] == '合欣' and c.get((i['order'], i['item'])) != i['qty']]
print('mismatch', bad)
